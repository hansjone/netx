"""Small xlsx/csv reader/writer for imports and templates (replaces pandas).

Rows come back as ``(row_no, {header: value})`` where ``row_no`` is the 1-based
sheet/file row (header = row 1), so import errors point at the row users see.
Blank rows are skipped. Excel cells keep their openpyxl types (str / int /
float / datetime / bool / None); integral floats become ints. CSV cells are
strings, empty cells are None.
"""
from __future__ import annotations

import csv
import io
import re
from collections.abc import Iterable, Iterator
from datetime import datetime
from typing import Any

from openpyxl import Workbook, load_workbook

Row = tuple[int, dict[str, Any]]


def _headers(raw: Iterable[Any]) -> list[str]:
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, h in enumerate(raw):
        name = "" if h is None else str(h).strip()
        if not name:
            name = f"Unnamed: {i}"
        n = seen.get(name, 0)
        seen[name] = n + 1
        out.append(name if n == 0 else f"{name}.{n}")
    return out


def _cell(v: Any) -> Any:
    if isinstance(v, str):
        return v if v.strip() else None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def _rows(records: Iterator[tuple[int, list[Any]]]) -> tuple[list[str], list[Row]]:
    headers: list[str] | None = None
    rows: list[Row] = []
    for row_no, values in records:
        cells = [_cell(v) for v in values]
        if all(c is None for c in cells):
            continue
        if headers is None:
            headers = _headers(values)
            continue
        cells += [None] * (len(headers) - len(cells))
        rows.append((row_no, dict(zip(headers, cells))))
    return headers or [], rows


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        # CSVs saved by Chinese Excel default to GBK.
        return content.decode("gb18030")


def read_csv(content: bytes) -> tuple[list[str], list[Row]]:
    reader = csv.reader(io.StringIO(_decode(content), newline=""))
    return _rows((i, rec) for i, rec in enumerate(reader, start=1))


def read_xlsx(content: bytes) -> tuple[list[str], list[Row]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        return _rows((i, list(rec)) for i, rec in enumerate(ws.iter_rows(values_only=True), start=1))
    finally:
        wb.close()


def read_table(content: bytes, filename: str = "") -> tuple[list[str], list[Row]]:
    if str(filename or "").lower().endswith(".csv"):
        return read_csv(content)
    return read_xlsx(content)


def write_csv(headers: list[str], rows: Iterable[Iterable[Any]]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(headers)
    for r in rows:
        writer.writerow(["" if v is None else v for v in r])
    # BOM so Excel opens UTF-8 (Chinese) correctly.
    return buf.getvalue().encode("utf-8-sig")


def write_xlsx(sheets: list[tuple[str, list[str], Iterable[Iterable[Any]]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for r in rows:
            ws.append(list(r))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%Y%m%d%H%M%S",
    "%Y%m%d",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y",
    "%d-%b-%Y %H:%M:%S",
    "%b %d, %Y %H:%M:%S",
)
_CN_DATE = re.compile(r"(\d{4})年(\d{1,2})月(\d{1,2})日")


def parse_datetime(v: Any) -> datetime | None:
    """Best-effort timestamp parse for spreadsheet cells; None when unparseable."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v
    text = str(v).strip()
    if not text or text.lower() in ("nan", "nat", "none", "null"):
        return None
    text = _CN_DATE.sub(r"\1-\2-\3", text)
    # 2024/1/2 03:04:05 -> 2024-1-2 ... ; zero-pad so fromisoformat/strptime accept it.
    if re.match(r"^\d{4}/\d{1,2}/\d{1,2}", text):
        text = text.replace("/", "-", 2)
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})(.*)$", text)
    if m:
        text = f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}{m.group(4)}"
    iso = text[:-1] + "+00:00" if text.endswith(("Z", "z")) else text
    try:
        return datetime.fromisoformat(iso)
    except ValueError:
        pass
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None
