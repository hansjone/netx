"""xlsx/csv helpers that replaced pandas in alarm + managed-NE imports."""
import io
from datetime import datetime, timezone

from openpyxl import Workbook, load_workbook

from netx_api.importer import _dt, _int, _str
from netx_api.tabular import parse_datetime, read_csv, read_table, read_xlsx, write_csv, write_xlsx


def _xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_xlsx_types_blank_rows_and_sheet_row_numbers():
    content = _xlsx([
        ["Time", "Code", "Port", None, "Code"],
        [datetime(2024, 1, 2, 3, 4, 5), 1001, 22.0, "x", "dup"],
        [None, None, None, None, None],
        ["  ", "A", 2.5, None, None],
    ])
    headers, rows = read_xlsx(content)
    assert headers == ["Time", "Code", "Port", "Unnamed: 3", "Code.1"]
    assert rows == [
        (2, {"Time": datetime(2024, 1, 2, 3, 4, 5), "Code": 1001, "Port": 22, "Unnamed: 3": "x", "Code.1": "dup"}),
        (4, {"Time": None, "Code": "A", "Port": 2.5, "Unnamed: 3": None, "Code.1": None}),
    ]


def test_read_csv_bom_quotes_short_rows_and_blank_lines():
    content = "﻿a,b,c\n1,\"x, y\",\n\n2\n".encode("utf-8")
    headers, rows = read_csv(content)
    assert headers == ["a", "b", "c"]
    assert rows == [(2, {"a": "1", "b": "x, y", "c": None}), (4, {"a": "2", "b": None, "c": None})]


def test_read_csv_falls_back_to_gbk():
    headers, rows = read_csv("名称,地址\n核心,10.0.0.1\n".encode("gbk"))
    assert headers == ["名称", "地址"]
    assert rows == [(2, {"名称": "核心", "地址": "10.0.0.1"})]


def test_read_table_dispatches_on_extension():
    assert read_table(b"h\nv\n", "X.CSV") == (["h"], [(2, {"h": "v"})])
    assert read_table(_xlsx([["h"], ["v"]]), "x.xlsx") == (["h"], [(2, {"h": "v"})])


def test_write_csv_and_xlsx_round_trip():
    csv_bytes = write_csv(["a", "b"], [[1, None], ["中文", "x"]])
    assert csv_bytes.startswith("﻿".encode("utf-8"))
    assert read_csv(csv_bytes)[1] == [(2, {"a": "1", "b": None}), (3, {"a": "中文", "b": "x"})]

    xlsx = write_xlsx([("import", ["a"], [[1]]), ("types", ["t"], [["x"], ["y"]])])
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["import", "types"]
    assert read_xlsx(xlsx) == (["a"], [(2, {"a": 1})])


def test_parse_datetime_formats():
    want = datetime(2024, 1, 2, 3, 4, 5)
    for text in (
        "2024-01-02 03:04:05", "2024/01/02 03:04:05", "2024/1/2 3:04:05", "2024-1-2 03:04:05",
        "2024-01-02T03:04:05", "01/02/2024 03:04:05", "20240102030405", "02-Jan-2024 03:04:05",
        "Jan 02, 2024 03:04:05", "2024年1月2日 03:04:05",
    ):
        assert parse_datetime(text) == want, text
    assert parse_datetime("2024-01-02T03:04:05Z") == datetime(2024, 1, 2, 3, 4, 5, tzinfo=timezone.utc)
    assert parse_datetime("2024-01-02") == datetime(2024, 1, 2)
    for bad in (None, "", "  ", "nan", "NaT", "abc"):
        assert parse_datetime(bad) is None


def test_importer_cell_helpers():
    assert _dt("2024-01-02T03:04:05Z") == datetime(2024, 1, 2, 3, 4, 5)  # stored as naive UTC
    assert _dt(float("nan")) is None and _dt("garbage") is None
    assert _str(1001) == "1001" and _str(None) == "" and _str(" nan ") == "" and _str(" a ") == "a"
    assert _int("7") == 7 and _int(2.0) == 2 and _int("x") == 0 and _int(None) == 0
